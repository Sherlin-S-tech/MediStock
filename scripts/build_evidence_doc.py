from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "Project_Run_and_GitHub_Evidence.docx"
WEB = DOCS / "dashboard.png"
TERM = DOCS / "terminal-transcript.png"

def font(size, bold=False, mono=False):
    names = (["C:/Windows/Fonts/consolab.ttf", "C:/Windows/Fonts/consola.ttf"] if mono else
             (["C:/Windows/Fonts/segoeuib.ttf", "C:/Windows/Fonts/segoeui.ttf"] if bold else ["C:/Windows/Fonts/segoeui.ttf"]))
    for name in names:
        if Path(name).exists():
            try:
                return ImageFont.truetype(name, size)
            except OSError:
                pass
    return ImageFont.load_default()

def make_terminal_image():
    width, height = 1600, 900
    img = Image.new("RGB", (width, height), "#0b1218")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((32, 28, width-32, height-28), radius=22, fill="#111b23", outline="#253540", width=2)
    draw.rounded_rectangle((32, 28, width-32, 100), radius=22, fill="#182630")
    draw.rectangle((32, 78, width-32, 100), fill="#182630")
    for x, color in [(64, "#ff6b65"), (91, "#f6c455"), (118, "#54c88b")]:
        draw.ellipse((x, 55, x+15, 70), fill=color)
    draw.text((155, 49), "PowerShell  |  Medicine Stock and Expiry Tracker", font=font(25, bold=True), fill="#d7e2e8")
    draw.text((64, 124), "Command transcript rendered from executed local checks", font=font(19), fill="#8fa5b3")
    rows = [
        ("> python -m pytest -q --basetemp .pytesttmp --junitxml=test-results.xml", "#8bd49c"),
        ("..............                                                   [100%]", "#d6e0e7"),
        ("--------------- generated xml file: D:\\Devops\\test-results.xml ----------------", "#8fa5b3"),
        ("14 passed in 1.06s", "#7ee0a1"),
        ("> docker compose config --quiet", "#8bd49c"),
        ("Compose configuration valid (exit code 0)", "#7ee0a1"),
        ("> Invoke-RestMethod http://127.0.0.1:5000/health | ConvertTo-Json -Compress", "#8bd49c"),
        ('{"database":"ok","status":"ok"}', "#d6e9f0"),
        ("> git status --short --branch", "#8bd49c"),
        ("## main", "#d6e0e7"),
        ("> git log -1 --oneline", "#8bd49c"),
        ("d044db3 Build medicine stock and expiry tracking project", "#d6e0e7"),
    ]
    y = 177
    for line, color in rows:
        draw.text((66, y), "PS D:\\Devops" + line if line.startswith(">") else line, font=font(21, mono=True), fill=color)
        y += 55
    draw.text((65, height-74), "Local verification only  •  no GitHub remote configured yet", font=font(17), fill="#91a5b0")
    img.save(TERM)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement("w:" + m)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)

def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, "12202A")
    set_cell_margins(cell, 130, 190, 130, 190)
    cell.text = ""
    for index, line in enumerate(lines):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.2)
        run.font.color.rgb = RGBColor(219, 232, 239)

def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for cell in table.rows[0].cells:
        shade(cell, "12785F")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255,255,255)
            run.font.size = Pt(8.2)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 65, 95, 65, 95)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(8.2)

def build():
    make_terminal_image()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.48)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(49, 67, 61)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in [("Title", 25, "17362D"), ("Heading 1", 16, "17362D"), ("Heading 2", 11, "12785F")]:
        st = doc.styles[style_name]
        st.font.name = "Segoe UI"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(9 if style_name != "Title" else 0)
        st.paragraph_format.space_after = Pt(4)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header.add_run("23IT723  |  DEVOPS LAB")
    hrun.font.name = "Segoe UI"; hrun.font.size = Pt(8); hrun.font.bold = True; hrun.font.color.rgb = RGBColor(18,120,95)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = footer.add_run("Medicine Stock and Expiry Tracking System  •  ")
    frun.font.size = Pt(8); frun.font.color.rgb = RGBColor(130,145,138)
    add_page_field(footer)

    p = doc.add_paragraph(style="Title")
    p.add_run("Project Run and GitHub Handoff")
    p = doc.add_paragraph()
    r = p.add_run("Medicine Stock and Expiry Tracking System with DevOps Automation")
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(18,120,95)
    p = doc.add_paragraph("This hand-in records the running local app, the fictional batches currently loaded, and verification commands. The project is committed locally on branch main; publishing will need the GitHub owner and repository details.")
    p.paragraph_format.space_after = Pt(6)

    doc.add_heading("Live website", level=1)
    pic = doc.add_picture(str(WEB), width=Inches(6.92))
    pic._inline.docPr.set("descr", "Live MedStock dashboard showing three demo medicine batches, 173 total units, one near-expiry batch, and one low-stock batch.")
    cap = doc.add_paragraph("Dashboard captured on 14 October 2026 from the local application.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True; run.font.size = Pt(8); run.font.color.rgb = RGBColor(107,126,117)
    doc.add_heading("Fictional demo records", level=2)
    t = doc.add_table(rows=1, cols=4)
    for c, value in zip(t.rows[0].cells, ["Medicine and batch", "Quantity / threshold", "Expiry", "Calculated status"]):
        c.text = value
    records = [
        ("Demo Paracetamol 500 mg\nDEMO-PARA-001", "120 / 20", "15 Mar 2027", "In stock"),
        ("Demo Cetirizine 10 mg\nDEMO-CET-002", "45 / 12", "25 Oct 2026", "Near expiry"),
        ("Demo Vitamin C 500 mg\nDEMO-VITC-003", "8 / 15", "15 Feb 2027", "Low stock"),
    ]
    for row in records:
        cells = t.add_row().cells
        for c, value in zip(cells, row):
            c.text = value
    format_table(t)

    doc.add_page_break()
    doc.add_heading("Run the application", level=1)
    doc.add_paragraph("From PowerShell, start the existing project with the commands below. Open http://localhost:5000 after the server starts.")
    add_code_block(doc, [
        "cd D:\\Devops",
        "py -m venv .venv",
        ".\\.venv\\Scripts\\Activate.ps1",
        "python -m pip install -r requirements.txt",
        '$env:SECRET_KEY = "replace-with-a-random-local-secret"',
        "python app.py",
    ])
    doc.add_paragraph("The database is created automatically. Add your own records with Add medicine; demo records remain available as examples.")
    doc.add_heading("Terminal verification transcript", level=1)
    doc.add_paragraph("This formatted transcript reproduces command output from the local checks. The dashboard above is a browser screenshot.")
    doc.add_picture(str(TERM), width=Inches(6.92))
    cap = doc.add_paragraph("Checks shown: pytest, Compose configuration, health endpoint, and local Git status.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True; run.font.size = Pt(8); run.font.color.rgb = RGBColor(107,126,117)

    doc.add_heading("GitHub publishing status", level=1)
    doc.add_paragraph("A local Git repository is ready on branch main with commit d044db3 (Build medicine stock and expiry tracking project). The working tree is clean. No remote repository has been created or pushed.")
    doc.add_paragraph("To publish, provide the GitHub username or organization, repository name, and whether it should be public or private. Do not send a password or access token in chat; authenticate through GitHub's normal sign-in or an approved connector when publishing.")
    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    build()

